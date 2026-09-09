from flask import Flask, render_template, request, redirect, url_for, send_file, g, jsonify, flash, session
import sqlite3
import pandas as pd
import io
import os
import sys
import shutil
import socket
import requests
from datetime import datetime, timedelta
from database import init_db, get_db_path, get_base_dir, extraer_litros_preciso, clasificar_tipo_producto, clasificar_viscosidad_tipo

SYNC_TOKEN = "oleos_minerales_sync_token_2026"

def get_resource_path(relative_path):
    """Obtiene la ruta absoluta a recursos internos (plantillas/estáticos) empaquetados con PyInstaller."""
    if hasattr(sys, '_MEIPASS'):
        return os.path.join(sys._MEIPASS, relative_path)
    return os.path.join(os.path.dirname(os.path.abspath(__file__)), relative_path)

app = Flask(
    __name__,
    template_folder=get_resource_path('templates'),
    static_folder=get_resource_path('static')
)
app.secret_key = 'oleos_minerales_secret_key_pyme'

def conectar():
    conn = sqlite3.connect(get_db_path(), timeout=10.0)
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("PRAGMA synchronous=NORMAL")
    conn.execute("PRAGMA busy_timeout=5000")
    conn.execute("PRAGMA foreign_keys=ON")
    return conn

def obtener_pin_admin():
    try:
        conn = conectar()
        cursor = conn.cursor()
        cursor.execute("SELECT pin_admin FROM configuracion_admin WHERE id = 1")
        row = cursor.fetchone()
        conn.close()
        if row and row[0]:
            return str(row[0])
    except Exception:
        pass
    return "1234"

def actualizar_pin_admin(nuevo_pin):
    try:
        conn = conectar()
        cursor = conn.cursor()
        cursor.execute("INSERT OR REPLACE INTO configuracion_admin (id, pin_admin) VALUES (1, ?)", (str(nuevo_pin).strip(),))
        conn.commit()
        conn.close()
        return True
    except Exception:
        return False

def obtener_ip_local():
    """Obtiene la IP local de la máquina en la red Wi-Fi/LAN."""
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.settimeout(0.2)
        s.connect(("8.8.8.8", 80))
        ip = s.getsockname()[0]
        s.close()
        return ip
    except Exception:
        try:
            return socket.gethostbyname(socket.gethostname())
        except Exception:
            return "127.0.0.1"

def obtener_perfil_actual():
    """Lee el perfil de la máquina desde la sesión o archivo perfil.txt."""
    if 'modo' in session:
        return session['modo']
    perfil_file = os.path.join(get_base_dir(), 'perfil.txt')
    if os.path.exists(perfil_file):
        try:
            with open(perfil_file, 'r', encoding='utf-8') as f:
                val = f.read().strip().lower()
                if 'mostrador' in val:
                    return 'mostrador'
                return 'admin'
        except Exception:
            pass
    return 'admin'

@app.route('/')
def index():
    modo_param = request.args.get('modo')
    if modo_param in ['admin', 'mostrador']:
        session['modo'] = modo_param

    perfil = obtener_perfil_actual()
    ip_local = obtener_ip_local()

    conn = conectar()
    cursor = conn.cursor()
    
    # 1. Cargar productos activos del mostrador junto con la suma de reservas pendientes
    cursor.execute('''
        SELECT 
            p.id, 
            p.codigo_barras, 
            p.nombre, 
            p.viscosidad, 
            p.presentacion, 
            p.bulto_unidades, 
            p.precio_costo, 
            p.precio_venta, 
            p.stock_actual, 
            p.ultima_venta,
            COALESCE(SUM(CASE WHEN r.estado = 'PENDIENTE' THEN r.cantidad ELSE 0 END), 0) AS reservado_pendiente
        FROM productos p
        LEFT JOIN reservas r ON p.id = r.producto_id
        WHERE p.ubicacion = 'MOSTRADOR' AND (p.activo = 1 OR p.activo IS NULL)
        GROUP BY p.id
    ''')
    productos = cursor.fetchall()

    # 2. Productos en Depósito 60D
    cursor.execute('''
        SELECT id, nombre, viscosidad, presentacion, bulto_unidades, precio_costo, precio_venta, stock_actual, ultima_venta 
        FROM productos 
        WHERE ubicacion = 'DEPOSITO_60D' AND (activo = 1 OR activo IS NULL)
    ''')
    productos_deposito = cursor.fetchall()

    # 3. Productos dados de baja (Papelera / Archivados con auto-limpieza a los 30 días)
    cursor.execute("""
        UPDATE historial_ventas 
        SET producto_id = NULL 
        WHERE producto_id IN (SELECT id FROM productos WHERE activo = 0 AND fecha_baja IS NOT NULL AND fecha_baja <= datetime('now', '-30 days'))
    """)
    cursor.execute("""
        DELETE FROM reservas 
        WHERE producto_id IN (SELECT id FROM productos WHERE activo = 0 AND fecha_baja IS NOT NULL AND fecha_baja <= datetime('now', '-30 days'))
    """)
    cursor.execute("DELETE FROM productos WHERE activo = 0 AND fecha_baja IS NOT NULL AND fecha_baja <= datetime('now', '-30 days')")
    conn.commit()

    cursor.execute("SELECT id, codigo_barras, nombre, viscosidad, fecha_baja FROM productos WHERE activo = 0 ORDER BY fecha_baja DESC")
    productos_archivados = cursor.fetchall()

    # 4. Reservas pendientes
    cursor.execute('''
        SELECT r.id, p.nombre, p.viscosidad, p.presentacion, r.empleado_nombre, r.cantidad, r.estado, r.fecha, p.stock_actual
        FROM reservas r
        JOIN productos p ON r.producto_id = p.id
        WHERE r.estado = 'PENDIENTE'
        ORDER BY r.fecha DESC
    ''')
    reservas_pendientes = cursor.fetchall()

    hace_60_dias = (datetime.now() - timedelta(days=60)).strftime('%Y-%m-%d %H:%M:%S')

    # 5. LISTA DE PERIODOS / MESES CON VENTAS Y FILTRADO
    cursor.execute("""
        SELECT DISTINCT strftime('%Y-%m', fecha) as mes 
        FROM historial_ventas 
        WHERE fecha IS NOT NULL 
        ORDER BY mes DESC
    """)
    meses_disponibles = [r[0] for r in cursor.fetchall() if r[0]]

    mes_actual_str = datetime.now().strftime('%Y-%m')
    mes_seleccionado = request.args.get('mes', '').strip()
    if not mes_seleccionado:
        if mes_actual_str in meses_disponibles:
            mes_seleccionado = mes_actual_str
        elif meses_disponibles:
            mes_seleccionado = meses_disponibles[0]
        else:
            mes_seleccionado = 'todos'

    if mes_seleccionado != 'todos' and mes_seleccionado:
        cursor.execute("""
            SELECT 
                COALESCE(SUM(CASE WHEN tipo_producto = 'LIQUIDO' THEN litros_totales ELSE 0 END), 0),
                COALESCE(SUM(CASE WHEN tipo_producto = 'GRASA' THEN litros_totales ELSE 0 END), 0),
                COALESCE(SUM(litros_totales), 0),
                COALESCE(SUM(precio_total), 0)
            FROM historial_ventas 
            WHERE strftime('%Y-%m', fecha) = ?
        """, (mes_seleccionado,))
        row_tot = cursor.fetchone()
        total_litros_liquidos = row_tot[0] or 0.0
        total_kg_grasas = row_tot[1] or 0.0
        total_litros = row_tot[2] or 0.0
        total_facturacion = row_tot[3] or 0.0

        cursor.execute('''
            SELECT 
                marca, 
                COALESCE(SUM(CASE WHEN tipo_producto = 'LIQUIDO' THEN litros_totales ELSE 0 END), 0) as litros_liquidos,
                COALESCE(SUM(CASE WHEN tipo_producto = 'GRASA' THEN litros_totales ELSE 0 END), 0) as kg_grasas,
                SUM(precio_total) as facturacion 
            FROM historial_ventas 
            WHERE strftime('%Y-%m', fecha) = ?
            GROUP BY marca 
            ORDER BY (litros_liquidos + kg_grasas) DESC
        ''', (mes_seleccionado,))
        metricas_marca = cursor.fetchall()

        cursor.execute('''
            SELECT h.tipo_producto, h.viscosidad, p.nombre, h.presentacion, h.cantidad, h.litros_totales
            FROM historial_ventas h
            LEFT JOIN productos p ON h.producto_id = p.id
            WHERE strftime('%Y-%m', h.fecha) = ?
        ''', (mes_seleccionado,))
        ventas_periodo = cursor.fetchall()
    else:
        cursor.execute("""
            SELECT 
                COALESCE(SUM(CASE WHEN tipo_producto = 'LIQUIDO' THEN litros_totales ELSE 0 END), 0),
                COALESCE(SUM(CASE WHEN tipo_producto = 'GRASA' THEN litros_totales ELSE 0 END), 0),
                COALESCE(SUM(litros_totales), 0),
                COALESCE(SUM(precio_total), 0)
            FROM historial_ventas
        """)
        row_tot = cursor.fetchone()
        total_litros_liquidos = row_tot[0] or 0.0
        total_kg_grasas = row_tot[1] or 0.0
        total_litros = row_tot[2] or 0.0
        total_facturacion = row_tot[3] or 0.0

        cursor.execute('''
            SELECT 
                marca, 
                COALESCE(SUM(CASE WHEN tipo_producto = 'LIQUIDO' THEN litros_totales ELSE 0 END), 0) as litros_liquidos,
                COALESCE(SUM(CASE WHEN tipo_producto = 'GRASA' THEN litros_totales ELSE 0 END), 0) as kg_grasas,
                SUM(precio_total) as facturacion 
            FROM historial_ventas 
            GROUP BY marca 
            ORDER BY (litros_liquidos + kg_grasas) DESC
        ''')
        metricas_marca = cursor.fetchall()

        cursor.execute('''
            SELECT h.tipo_producto, h.viscosidad, p.nombre, h.presentacion, h.cantidad, h.litros_totales
            FROM historial_ventas h
            LEFT JOIN productos p ON h.producto_id = p.id
        ''')
        ventas_periodo = cursor.fetchall()

    # Procesar inteligentemente viscosidades y tipos de grasa para la cuadrícula
    resumen_visc = {}
    for item in ventas_periodo:
        tipo_p, visc, nom_p, pres_p, cant, vol = item
        etiqueta, categoria = clasificar_viscosidad_tipo(tipo_p, visc, nom_p, pres_p)
        if etiqueta not in resumen_visc:
            resumen_visc[etiqueta] = {
                'etiqueta': etiqueta,
                'volumen': 0.0,
                'unidades': 0,
                'unidad': 'Kg' if categoria == 'GRASA' else 'Lts',
                'es_grasa': (categoria == 'GRASA')
            }
        resumen_visc[etiqueta]['volumen'] += (vol or 0.0)
        resumen_visc[etiqueta]['unidades'] += (cant or 0)

    metricas_viscosidad = sorted(
        [(v['etiqueta'], v['volumen'], v['unidades'], v['unidad'], v['es_grasa']) for v in resumen_visc.values()],
        key=lambda x: x[1],
        reverse=True
    )

    cursor.execute('''
        SELECT 
            strftime('%Y-%m', fecha) as mes,
            COUNT(id) as operaciones,
            COALESCE(SUM(CASE WHEN tipo_producto = 'LIQUIDO' THEN litros_totales ELSE 0 END), 0) as litros_liquidos,
            COALESCE(SUM(CASE WHEN tipo_producto = 'GRASA' THEN litros_totales ELSE 0 END), 0) as kg_grasas,
            COALESCE(SUM(precio_total), 0) as facturacion,
            COALESCE(SUM(litros_totales), 0) as total_volumen
        FROM historial_ventas
        GROUP BY strftime('%Y-%m', fecha)
        ORDER BY mes DESC
    ''')
    resumen_mensual = cursor.fetchall()

    # 5b. HISTORIAL DE REMITOS Y VENTAS (Para visualización y control en Gestión)
    cursor.execute('''
        SELECT 
            h.fecha,
            COUNT(h.id) as total_items,
            SUM(h.litros_totales) as litros_remito,
            SUM(h.precio_total) as monto_remito,
            strftime('%Y-%m', h.fecha) as mes_remito,
            GROUP_CONCAT(DISTINCT h.marca) as marcas_resumen,
            COALESCE(SUM(CASE WHEN h.tipo_producto = 'LIQUIDO' THEN h.litros_totales ELSE 0 END), 0) as litros_liquidos,
            COALESCE(SUM(CASE WHEN h.tipo_producto = 'GRASA' THEN h.litros_totales ELSE 0 END), 0) as kg_grasas,
            GROUP_CONCAT(COALESCE(h.nombre_producto, p.nombre, h.marca) || ' (' || h.cantidad || ' u.)', ' | ') as detalle_items
        FROM historial_ventas h
        LEFT JOIN productos p ON h.producto_id = p.id
        GROUP BY h.fecha
        ORDER BY h.fecha DESC
    ''')
    historial_remitos = cursor.fetchall()

    # 6. CÁLCULO INTELIGENTE DE STOCK BAJO, AGOTADOS Y LISTA DE REPOSICIÓN
    stock_bajo_count = 0
    agotados_count = 0
    productos_faltantes = []

    for p in productos:
        # p: (id, codigo_barras, nombre, viscosidad, presentacion, bulto_unidades, precio_costo, precio_venta, stock_actual, ultima_venta, reservado_pendiente)
        stock = p[8] if p[8] is not None else 0
        presentacion = str(p[4]).upper() if p[4] else ""
        bulto = p[5] if p[5] else 1
        es_individual = any(x in presentacion for x in ['BALDE', 'TAMBOR', 'TBR', 'CONTENEDOR', 'GRANEL'])
        
        if stock == 0:
            agotados_count += 1
            productos_faltantes.append({
                'id': p[0], 'codigo': p[1], 'nombre': p[2], 'viscosidad': p[3],
                'presentacion': p[4], 'bulto': bulto, 'es_individual': es_individual,
                'precio_costo': p[6], 'precio_venta': p[7], 'stock': 0, 'estado': 'AGOTADO'
            })
        else:
            umbral = 1 if es_individual else (bulto if bulto > 1 else 5)
            if stock <= umbral:
                stock_bajo_count += 1
                productos_faltantes.append({
                    'id': p[0], 'codigo': p[1], 'nombre': p[2], 'viscosidad': p[3],
                    'presentacion': p[4], 'bulto': bulto, 'es_individual': es_individual,
                    'precio_costo': p[6], 'precio_venta': p[7], 'stock': stock, 'estado': 'BAJO_STOCK'
                })

    conn.close()

    return render_template(
        'index.html', 
        productos=productos, 
        productos_deposito=productos_deposito,
        productos_archivados=productos_archivados,
        reservas_pendientes=reservas_pendientes,
        hace_60_dias=hace_60_dias,
        total_litros=total_litros,
        total_litros_liquidos=total_litros_liquidos,
        total_kg_grasas=total_kg_grasas,
        total_facturacion=total_facturacion,
        meses_disponibles=meses_disponibles,
        mes_seleccionado=mes_seleccionado,
        metricas_marca=metricas_marca,
        metricas_viscosidad=metricas_viscosidad,
        resumen_mensual=resumen_mensual,
        stock_bajo_count=stock_bajo_count,
        agotados_count=agotados_count,
        productos_faltantes=productos_faltantes,
        historial_remitos=historial_remitos,
        perfil=perfil,
        ip_local=ip_local,
        pin_admin=obtener_pin_admin()
    )

@app.route('/crear_reserva', methods=['POST'])
def crear_reserva():
    id_prod = request.form['id_producto']
    vendedor = request.form['vendedor']
    cantidad = int(request.form['cantidad'])

    conn = conectar()
    cursor = conn.cursor()
    cursor.execute("SELECT stock_actual, nombre FROM productos WHERE id = ?", (id_prod,))
    prod = cursor.fetchone()

    if prod:
        stock_actual = prod[0]
        nombre = prod[1]
        if stock_actual >= cantidad:
            cursor.execute('''
                INSERT INTO reservas (producto_id, empleado_nombre, cantidad, estado)
                VALUES (?, ?, ?, 'PENDIENTE')
            ''', (id_prod, vendedor, cantidad))
            conn.commit()
            flash(f"✅ Solicitud de reserva registrada para {nombre} ({cantidad} u.).", "success")
        else:
            flash(f"❌ No se puede reservar: stock insuficiente para {nombre} (Solicitado: {cantidad} u. | Disponible: {stock_actual} u.).", "error")
    else:
        flash("❌ Error: Producto no encontrado.", "error")

    conn.close()
    return redirect(url_for('index'))

@app.route('/aprobar_reserva/<int:id>')
def aprobar_reserva(id):
    conn = conectar()
    cursor = conn.cursor()
    cursor.execute("SELECT producto_id, cantidad FROM reservas WHERE id = ?", (id,))
    res = cursor.fetchone()
    
    if res:
        id_prod, cantidad = res
        cursor.execute("SELECT nombre, viscosidad, presentacion, precio_venta, stock_actual, litros_unitarios, tipo_producto FROM productos WHERE id = ?", (id_prod,))
        prod = cursor.fetchone()
        
        if prod and prod[4] >= cantidad:
            nombre, viscosidad, presentacion, precio_venta, stock_actual, litros_unitarios, tipo_prod = prod
            marca = nombre.split(' - ')[0] if ' - ' in nombre else 'OTRA'
            nuevo_stock = stock_actual - cantidad
            litros_vendidos = cantidad * (litros_unitarios or 1.0)
            precio_total = cantidad * precio_venta
            fecha_actual = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
            tipo_final = tipo_prod if tipo_prod else clasificar_tipo_producto(nombre, presentacion)

            cursor.execute("UPDATE productos SET stock_actual = ?, ultima_venta = ? WHERE id = ?", (nuevo_stock, fecha_actual, id_prod))
            cursor.execute("UPDATE reservas SET estado = 'APROBADO' WHERE id = ?", (id,))
            cursor.execute('''
                INSERT INTO historial_ventas (producto_id, marca, viscosidad, presentacion, cantidad, litros_totales, precio_total, fecha, tipo_producto)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            ''', (id_prod, marca, viscosidad, presentacion, cantidad, litros_vendidos, precio_total, fecha_actual, tipo_final))

            conn.commit()

    conn.close()
    return redirect(url_for('index'))

@app.route('/rechazar_reserva/<int:id>')
def rechazar_reserva(id):
    conn = conectar()
    cursor = conn.cursor()
    cursor.execute("UPDATE reservas SET estado = 'RECHAZADO' WHERE id = ?", (id,))
    conn.commit()
    conn.close()
    return redirect(url_for('index'))

@app.route('/editar_stock', methods=['POST'])
def editar_stock():
    id_prod = request.form['id_producto']
    nuevo_stock = int(request.form['nuevo_stock'])
    
    conn = conectar()
    cursor = conn.cursor()
    cursor.execute("UPDATE productos SET stock_actual = ? WHERE id = ?", (nuevo_stock, id_prod))
    conn.commit()
    conn.close()
    return redirect(url_for('index'))

@app.route('/agregar', methods=['POST'])
def agregar():
    marca = request.form['marca'].strip()
    nombre_linea = request.form['nombre'].strip()
    nombre_completo = f"{marca} - {nombre_linea}" if marca else nombre_linea
    viscosidad = request.form['viscosidad'].strip()
    presentacion = request.form['presentacion'].strip()
    bulto = int(request.form.get('bulto', 1))
    codigo_barras_raw = request.form.get('codigo_barras', '').strip()
    # Si el código de barras está vacío, guardarlo como NULL para no violar UNIQUE
    codigo_barras = codigo_barras_raw if codigo_barras_raw else None
    precio_costo_raw = str(request.form.get('precio_costo', '0') or '0').strip().replace(',', '.')
    precio_venta_raw = str(request.form.get('precio_venta', '0') or '0').strip().replace(',', '.')
    try:
        precio_costo_unit = float(precio_costo_raw) if precio_costo_raw else 0.0
    except ValueError:
        precio_costo_unit = 0.0
    try:
        precio_venta_unit = float(precio_venta_raw) if precio_venta_raw else 0.0
    except ValueError:
        precio_venta_unit = 0.0
    precio_costo = precio_costo_unit * bulto
    precio_venta = precio_venta_unit * bulto
    stock_cajas = int(request.form.get('stock_cajas', 0) or 0)
    stock_unidades = int(request.form.get('stock_unidades', 0) or 0)
    stock = (stock_cajas * bulto) + stock_unidades

    litros_unitarios = extraer_litros_preciso(presentacion + ' ' + nombre_completo)
    tipo_producto = clasificar_tipo_producto(nombre_completo, presentacion)

    conn = conectar()
    cursor = conn.cursor()
    try:
        cursor.execute('''
            INSERT INTO productos (codigo_barras, nombre, viscosidad, presentacion, bulto_unidades, litros_unitarios, precio_costo, precio_venta, stock_actual, ubicacion, activo, tipo_producto)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 'MOSTRADOR', 1, ?)
        ''', (codigo_barras, nombre_completo, viscosidad, presentacion, bulto, litros_unitarios, precio_costo, precio_venta, stock, tipo_producto))
        conn.commit()
        flash(f"✅ Producto '{nombre_completo}' guardado correctamente.", "success")
    except Exception as e:
        conn.rollback()
        if 'UNIQUE' in str(e):
            flash(f"❌ Error: El código de barras '{codigo_barras}' ya está asignado a otro producto. Usá uno diferente o dejá el campo vacío.", "error")
        else:
            flash(f"❌ Error al guardar el producto: {str(e)}", "error")
    finally:
        conn.close()
    return redirect(url_for('index'))

@app.route('/vender', methods=['POST'])
def vender():
    id_prod = request.form['id_producto']
    cantidad = int(request.form['cantidad'])
    
    conn = conectar()
    cursor = conn.cursor()
    cursor.execute("SELECT nombre, viscosidad, presentacion, precio_venta, stock_actual, litros_unitarios, tipo_producto FROM productos WHERE id = ?", (id_prod,))
    prod = cursor.fetchone()
    
    if prod:
        nombre, viscosidad, presentacion, precio_venta, stock_actual, litros_unitarios, tipo_prod = prod
        if stock_actual >= cantidad:
            marca = nombre.split(' - ')[0] if ' - ' in nombre else 'OTRA'
            nuevo_stock = stock_actual - cantidad
            litros_vendidos = cantidad * (litros_unitarios or 1.0)
            precio_total = cantidad * precio_venta
            fecha_actual = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
            tipo_final = tipo_prod if tipo_prod else clasificar_tipo_producto(nombre, presentacion)

            cursor.execute("UPDATE productos SET stock_actual = ?, ultima_venta = ? WHERE id = ?", (nuevo_stock, fecha_actual, id_prod))
            cursor.execute('''
                INSERT INTO historial_ventas (producto_id, marca, viscosidad, presentacion, cantidad, litros_totales, precio_total, fecha, tipo_producto)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            ''', (id_prod, marca, viscosidad, presentacion, cantidad, litros_vendidos, precio_total, fecha_actual, tipo_final))

            conn.commit()
            flash(f"✅ Venta registrada: {cantidad} u. de {nombre}.", "success")
        else:
            flash(f"❌ No se puede vender: stock insuficiente para {nombre} (Solicitado: {cantidad} u. | Disponible: {stock_actual} u.).", "error")
    else:
        flash("❌ Error: Producto no encontrado.", "error")
    
    conn.close()
    return redirect(url_for('index'))

@app.route('/vender_remito', methods=['POST'])
def vender_remito():
    """Procesa una venta con múltiples productos (remito/listado)."""
    try:
        data = request.get_json()
        items = data.get('items', [])
        if not items:
            return jsonify({'success': False, 'message': 'El carrito está vacío.'}), 400

        conn = conectar()
        cursor = conn.cursor()
        fecha_remito_custom = (data.get('fecha_remito') or '').strip()
        if fecha_remito_custom:
            if len(fecha_remito_custom) == 10:
                fecha_actual = f"{fecha_remito_custom} {datetime.now().strftime('%H:%M:%S')}"
            else:
                fecha_actual = fecha_remito_custom
        else:
            fecha_actual = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        resultados = []

        for item in items:
            id_prod = item.get('id_producto')
            cantidad = int(item.get('cantidad', 1))
            cantidad_bonificada = int(item.get('cantidad_bonificada', 0))
            tipo_venta = item.get('tipo_venta', 'unidad')
            bulto_unidades_carrito = int(item.get('bulto_unidades', 1))

            cursor.execute(
                "SELECT nombre, viscosidad, presentacion, precio_venta, stock_actual, litros_unitarios, bulto_unidades, tipo_producto FROM productos WHERE id = ?",
                (id_prod,)
            )
            prod = cursor.fetchone()
            if not prod:
                conn.rollback()
                conn.close()
                return jsonify({'success': False, 'message': f'Producto ID {id_prod} no encontrado.'}), 404

            nombre, viscosidad, presentacion, precio_venta, stock_actual, litros_unitarios, bulto_unidades_db, tipo_prod = prod
            tipo_final = tipo_prod if tipo_prod else clasificar_tipo_producto(nombre, presentacion)
            
            # Detectar si es balde, tambor o granel para tratarlo como unidad individual (bulto = 1)
            pres_upper = (presentacion or '').upper()
            es_individual = any(w in pres_upper for w in ['BALDE', 'TAMBOR', 'TBR', 'CONTENEDOR', 'GRANEL'])
            bulto_unidades = 1 if es_individual else (bulto_unidades_db if bulto_unidades_db else bulto_unidades_carrito)

            # Calcular total a descontar de stock (en botellas individuales)
            if tipo_venta == 'bulto':
                total_a_descontar = (cantidad + cantidad_bonificada) * bulto_unidades
                litros_vendidos = (cantidad + cantidad_bonificada) * bulto_unidades * (litros_unitarios or 1.0)
                precio_total = cantidad * precio_venta  # precio_venta en DB es precio de caja
                cantidad_historial = (cantidad + cantidad_bonificada) * bulto_unidades
            else:
                total_a_descontar = cantidad + cantidad_bonificada
                litros_vendidos = (cantidad + cantidad_bonificada) * (litros_unitarios or 1.0)
                precio_unitario = precio_venta / bulto_unidades
                precio_total = cantidad * precio_unitario
                cantidad_historial = cantidad + cantidad_bonificada

            if stock_actual < total_a_descontar:
                conn.rollback()
                conn.close()
                return jsonify({'success': False, 'message': f'Stock insuficiente para "{nombre}": disponible {stock_actual} u., solicitado {total_a_descontar} u.'}), 400

            marca = nombre.split(' - ')[0] if ' - ' in nombre else 'OTRA'
            nuevo_stock = stock_actual - total_a_descontar

            cursor.execute("UPDATE productos SET stock_actual = ?, ultima_venta = ? WHERE id = ?", (nuevo_stock, fecha_actual, id_prod))
            cursor.execute('''
                INSERT INTO historial_ventas (producto_id, nombre_producto, marca, viscosidad, presentacion, cantidad, litros_totales, precio_total, fecha, tipo_producto)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ''', (id_prod, nombre, marca, viscosidad, presentacion, cantidad_historial, litros_vendidos, precio_total, fecha_actual, tipo_final))

            resultados.append({'nombre': nombre, 'cantidad': cantidad, 'precio_total': precio_total, 'tipo_venta': tipo_venta})

        conn.commit()
        conn.close()
        return jsonify({'success': True, 'message': f'Remito procesado: {len(resultados)} producto(s) descontados.', 'items': resultados})

    except Exception as e:
        return jsonify({'success': False, 'message': str(e)}), 500

@app.route('/editar_precios_barras', methods=['POST'])
def editar_precios_barras():
    id_producto = request.form.get('id_producto')
    codigo_barras = request.form.get('codigo_barras', '').strip()
    precio_costo_raw = str(request.form.get('precio_costo', '0') or '0').strip().replace(',', '.')
    precio_venta_raw = str(request.form.get('precio_venta', '0') or '0').strip().replace(',', '.')

    try:
        precio_costo_unit = float(precio_costo_raw) if precio_costo_raw else 0.0
    except ValueError:
        precio_costo_unit = 0.0
    try:
        precio_venta_unit = float(precio_venta_raw) if precio_venta_raw else 0.0
    except ValueError:
        precio_venta_unit = 0.0

    try:
        conn = conectar()
        cursor = conn.cursor()
        
        # Obtener bulto_unidades para poder calcular el precio de caja
        cursor.execute("SELECT bulto_unidades FROM productos WHERE id = ?", (id_producto,))
        row = cursor.fetchone()
        bulto = row[0] if row else 1
        
        precio_costo = precio_costo_unit * bulto
        precio_venta = precio_venta_unit * bulto
        
        cursor.execute("""
            UPDATE productos 
            SET codigo_barras = ?, precio_costo = ?, precio_venta = ? 
            WHERE id = ?
        """, (codigo_barras if codigo_barras else None, precio_costo, precio_venta, id_producto))
        
        conn.commit()
        conn.close()
        flash("✅ Producto actualizado correctamente.", "success")
    except Exception as e:
        flash(f"❌ Error al actualizar producto: {str(e)}", "error")

    return redirect(url_for('index'))

@app.route('/editar_precios_barras_ajax', methods=['POST'])
def editar_precios_barras_ajax():
    try:
        data = request.get_json()
        id_producto = data.get('id_producto')
        codigo_barras = data.get('codigo_barras', '').strip()
        precio_costo_raw = str(data.get('precio_costo', '0') or '0').strip().replace(',', '.')
        precio_venta_raw = str(data.get('precio_venta', '0') or '0').strip().replace(',', '.')
        try:
            precio_costo_unit = float(precio_costo_raw) if precio_costo_raw else 0.0
        except ValueError:
            precio_costo_unit = 0.0
        try:
            precio_venta_unit = float(precio_venta_raw) if precio_venta_raw else 0.0
        except ValueError:
            precio_venta_unit = 0.0
        # Nuevos campos de edición completa
        marca = data.get('marca', '').strip()
        nombre_linea = data.get('nombre_linea', '').strip()
        viscosidad = data.get('viscosidad', '').strip()
        presentacion = data.get('presentacion', '').strip()
        bulto_unidades = int(data.get('bulto_unidades', 1))

        # Convertir a precio de caja
        precio_costo = precio_costo_unit * bulto_unidades
        precio_venta = precio_venta_unit * bulto_unidades

        # Armar nombre completo: "MARCA - Línea" o solo la línea si no hay marca
        nombre_completo = f"{marca} - {nombre_linea}" if marca else nombre_linea
        litros_unitarios = extraer_litros_preciso(presentacion + ' ' + (nombre_completo or ''))

        conn = conectar()
        cursor = conn.cursor()
        cursor.execute("""
            UPDATE productos 
            SET codigo_barras = ?, precio_costo = ?, precio_venta = ?,
                nombre = ?, viscosidad = ?, presentacion = ?, bulto_unidades = ?, litros_unitarios = ?
            WHERE id = ?
        """, (
            codigo_barras if codigo_barras else None,
            precio_costo, precio_venta,
            nombre_completo if nombre_completo.strip(' -') else None,
            viscosidad, presentacion, bulto_unidades, litros_unitarios,
            id_producto
        ))
        conn.commit()
        conn.close()

        return jsonify({
            'success': True, 
            'message': 'Producto actualizado', 
            'nombre_completo': nombre_completo,
            'precio_venta': precio_venta,
            'precio_costo': precio_costo
        })
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)}), 500

@app.route('/mover_deposito/<int:id>')
def mover_deposito(id):
    conn = conectar()
    cursor = conn.cursor()
    cursor.execute("UPDATE productos SET ubicacion = 'DEPOSITO_60D' WHERE id = ?", (id,))
    conn.commit()
    conn.close()
    return redirect(url_for('index'))

@app.route('/mover_mostrador/<int:id>')
def mover_mostrador(id):
    conn = conectar()
    cursor = conn.cursor()
    cursor.execute("UPDATE productos SET ubicacion = 'MOSTRADOR', ultima_venta = CURRENT_TIMESTAMP WHERE id = ?", (id,))
    conn.commit()
    conn.close()
    return redirect(url_for('index'))

@app.route('/exportar')
def exportar():
    conn = conectar()
    df = pd.read_sql_query("SELECT id AS ID, codigo_barras AS Codigo, nombre AS Producto, viscosidad AS Viscosidad, presentacion AS Envase, bulto_unidades AS Bulto_Unidades, precio_costo AS Costo_Neto, precio_venta AS Venta_Final, stock_actual AS Stock FROM productos WHERE activo = 1 OR activo IS NULL", conn)
    conn.close()

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Inventario')
    output.seek(0)

    return send_file(output, download_name='Inventario_Oleos_Minerales.xlsx', as_attachment=True)

@app.route('/exportar_reposicion')
def exportar_reposicion():
    conn = conectar()
    cursor = conn.cursor()
    cursor.execute("""
        SELECT id, codigo_barras, nombre, viscosidad, presentacion, bulto_unidades, precio_costo, precio_venta, stock_actual
        FROM productos
        WHERE ubicacion = 'MOSTRADOR' AND (activo = 1 OR activo IS NULL)
        ORDER BY nombre ASC
    """)
    prods = cursor.fetchall()
    conn.close()
    
    filas = []
    for p in prods:
        stock = p[8] if p[8] is not None else 0
        pres = str(p[4]).upper() if p[4] else ""
        bulto = p[5] if p[5] else 1
        es_ind = any(x in pres for x in ['BALDE', 'TAMBOR', 'TBR', 'CONTENEDOR', 'GRANEL'])
        umbral = 1 if es_ind else (bulto if bulto > 1 else 5)
        
        if stock == 0 or stock <= umbral:
            estado = "⛔ AGOTADO (0 u.)" if stock == 0 else f"⚠️ BAJO STOCK ({stock} u.)"
            filas.append({
                'ID': p[0],
                'Codigo': p[1] or '',
                'Producto': p[2],
                'Viscosidad': p[3] or '',
                'Envase': p[4] or '',
                'Bulto_Unidades': bulto,
                'Stock_Actual': stock,
                'Estado_Alerta': estado,
                'Costo_Neto': p[6],
                'Precio_Venta': p[7]
            })
            
    df = pd.DataFrame(filas)
    if df.empty:
        df = pd.DataFrame(columns=['ID', 'Codigo', 'Producto', 'Viscosidad', 'Envase', 'Bulto_Unidades', 'Stock_Actual', 'Estado_Alerta', 'Costo_Neto', 'Precio_Venta'])
        
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Faltantes_Reposicion')
    output.seek(0)
    
    fecha_hoy = datetime.now().strftime("%Y-%m-%d")
    return send_file(
        output,
        download_name=f'Pedido_Reposicion_Faltantes_{fecha_hoy}.xlsx',
        as_attachment=True,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

@app.route('/exportar_cierre_mes/<mes>')
def exportar_cierre_mes(mes):
    """Genera una planilla Excel oficial con el cierre mensual de ventas y litros."""
    try:
        conn = conectar()
        cursor = conn.cursor()
        
        cursor.execute("""
            SELECT 
                COUNT(id), 
                COALESCE(SUM(CASE WHEN tipo_producto = 'LIQUIDO' THEN litros_totales ELSE 0 END), 0),
                COALESCE(SUM(CASE WHEN tipo_producto = 'GRASA' THEN litros_totales ELSE 0 END), 0),
                COALESCE(SUM(litros_totales), 0),
                COALESCE(SUM(precio_total), 0)
            FROM historial_ventas
            WHERE strftime('%Y-%m', fecha) = ?
        """, (mes,))
        tot_ops, tot_liq, tot_grasa, tot_vol, tot_pesos = cursor.fetchone()
        tot_ops = tot_ops or 0
        tot_liq = tot_liq or 0.0
        tot_grasa = tot_grasa or 0.0
        tot_vol = tot_vol or 0.0
        tot_pesos = tot_pesos or 0.0

        df_marcas = pd.read_sql_query("""
            SELECT 
                marca AS Marca, 
                ROUND(COALESCE(SUM(CASE WHEN tipo_producto = 'LIQUIDO' THEN litros_totales ELSE 0 END), 0), 2) AS Litros_Liquidos,
                ROUND(COALESCE(SUM(CASE WHEN tipo_producto = 'GRASA' THEN litros_totales ELSE 0 END), 0), 2) AS Kg_Grasas,
                ROUND(SUM(precio_total), 2) AS Facturacion_Total
            FROM historial_ventas
            WHERE strftime('%Y-%m', fecha) = ?
            GROUP BY marca
            ORDER BY (Litros_Liquidos + Kg_Grasas) DESC
        """, conn, params=(mes,))

        cursor.execute("""
            SELECT h.tipo_producto, h.viscosidad, p.nombre, h.presentacion, h.cantidad, h.litros_totales
            FROM historial_ventas h
            LEFT JOIN productos p ON h.producto_id = p.id
            WHERE strftime('%Y-%m', h.fecha) = ?
        """, (mes,))
        ventas_mes_visc = cursor.fetchall()
        res_visc = {}
        for row in ventas_mes_visc:
            tp, vi, np, pr, cant, vol = row
            etq, cat = clasificar_viscosidad_tipo(tp, vi, np, pr)
            if etq not in res_visc:
                res_visc[etq] = {'etiqueta': etq, 'unidades': 0, 'volumen': 0.0, 'unidad': 'Kg' if cat == 'GRASA' else 'Lts'}
            res_visc[etq]['unidades'] += (cant or 0)
            res_visc[etq]['volumen'] += (vol or 0.0)

        lista_visc = sorted(res_visc.values(), key=lambda x: x['volumen'], reverse=True)

        df_detalle = pd.read_sql_query("""
            SELECT fecha AS Fecha_Hora,
                   tipo_producto AS Tipo,
                   marca AS Marca,
                   viscosidad AS Viscosidad,
                   presentacion AS Envase,
                   cantidad AS Cantidad,
                   ROUND(litros_totales, 2) AS Litros_Kg,
                   ROUND(precio_total, 2) AS Subtotal_Pesos
            FROM historial_ventas
            WHERE strftime('%Y-%m', fecha) = ?
            ORDER BY fecha ASC
        """, conn, params=(mes,))
        conn.close()

        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment

        wb = openpyxl.Workbook()
        ws_resumen = wb.active
        ws_resumen.title = "Resumen Cierre"

        header_font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        
        ws_resumen.merge_cells("A1:D1")
        ws_resumen["A1"] = f"ÓLEOS MINERALES SRL - CIERRE DE MES ({mes})"
        ws_resumen["A1"].font = header_font
        ws_resumen["A1"].fill = header_fill
        ws_resumen["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws_resumen.row_dimensions[1].height = 28

        ws_resumen["A3"] = "Periodo de Cierre:"
        ws_resumen["B3"] = mes
        ws_resumen["A4"] = "Total Operaciones/Ventas:"
        ws_resumen["B4"] = tot_ops
        ws_resumen["A5"] = "Aceites y Líquidos (Litros):"
        ws_resumen["B5"] = round(tot_liq, 2)
        ws_resumen["A6"] = "Grasas Lubricantes (Kilos):"
        ws_resumen["B6"] = round(tot_grasa, 2)
        ws_resumen["A7"] = "Volumen Consolidado (Lts/Kg):"
        ws_resumen["B7"] = round(tot_vol, 2)
        ws_resumen["A8"] = "Total Facturación ($):"
        ws_resumen["B8"] = round(tot_pesos, 2)

        for r in range(3, 9):
            ws_resumen[f"A{r}"].font = Font(bold=True)
            ws_resumen[f"B{r}"].font = Font(bold=True, color="1F4E79")

        # Tabla Marcas
        ws_resumen["A10"] = "DESGLOSE POR MARCA"
        ws_resumen["A10"].font = Font(bold=True, size=11, color="1F4E79")
        ws_resumen.append([])
        ws_resumen.append(["Marca", "Aceites / Líquidos (Lts)", "Grasas (Kg)", "Facturación ($)"])
        r_marca = ws_resumen.max_row
        for col in [f"A{r_marca}", f"B{r_marca}", f"C{r_marca}", f"D{r_marca}"]:
            ws_resumen[col].font = Font(bold=True, color="FFFFFF")
            ws_resumen[col].fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")

        for _, row in df_marcas.iterrows():
            ws_resumen.append([row["Marca"], row["Litros_Liquidos"], row["Kg_Grasas"], row["Facturacion_Total"]])

        # Tabla Viscosidades y Tipos de Grasa
        ws_resumen.append([])
        r_visc_title = ws_resumen.max_row + 1
        ws_resumen[f"A{r_visc_title}"] = "DESGLOSE POR VISCOSIDAD Y TIPO DE GRASA"
        ws_resumen[f"A{r_visc_title}"].font = Font(bold=True, size=11, color="1F4E79")
        ws_resumen.append([])
        ws_resumen.append(["Viscosidad / Tipo de Grasa", "Unidades Vendidas", "Volumen Salido", "Unidad"])
        r_visc_hdr = ws_resumen.max_row
        for col in [f"A{r_visc_hdr}", f"B{r_visc_hdr}", f"C{r_visc_hdr}", f"D{r_visc_hdr}"]:
            ws_resumen[col].font = Font(bold=True, color="FFFFFF")
            ws_resumen[col].fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")

        for lv in lista_visc:
            ws_resumen.append([lv["etiqueta"], lv["unidades"], round(lv["volumen"], 2), lv["unidad"]])

        # Sheet 2: Detalle de Ventas
        ws_det = wb.create_sheet(title="Detalle de Ventas")
        ws_det.append(["Fecha y Hora", "Tipo (Líquido/Grasa)", "Marca", "Viscosidad", "Presentación / Envase", "Cantidad", "Litros / Kg", "Subtotal ($)"])
        for col_idx in range(1, 9):
            cell = ws_det.cell(row=1, column=col_idx)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill

        for _, row in df_detalle.iterrows():
            ws_det.append([row["Fecha_Hora"], row["Tipo"], row["Marca"], row["Viscosidad"], row["Envase"], row["Cantidad"], row["Litros_Kg"], row["Subtotal_Pesos"]])

        for ws in [ws_resumen, ws_det]:
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return send_file(
            output, 
            download_name=f'Cierre_Mensual_{mes}_Oleos_Minerales.xlsx', 
            as_attachment=True,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
    except Exception as e:
        flash(f"❌ Error al exportar cierre del mes: {str(e)}", "error")
        return redirect(url_for('index'))

@app.route('/iniciar_inventario_auditoria', methods=['POST'])
def iniciar_inventario_auditoria():
    conn = conectar()
    cursor = conn.cursor()
    
    # 1. Crear tabla de auditoría si no existe
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS auditoria_stock (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            fecha TEXT,
            id_producto INTEGER,
            stock_teorico INTEGER,
            stock_contado INTEGER DEFAULT 0
        )
    ''')
    
    # 2. Guardar foto del stock teórico antes del blanqueo
    fecha_actual = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    cursor.execute("SELECT id, stock_actual FROM productos WHERE activo = 1 OR activo IS NULL")
    productos = cursor.fetchall()
    
    for p in productos:
        cursor.execute("""
            INSERT INTO auditoria_stock (fecha, id_producto, stock_teorico, stock_contado)
            VALUES (?, ?, ?, 0)
        """, (fecha_actual, p[0], p[1]))
    
    # 3. Blanquear stock a cero
    cursor.execute("UPDATE productos SET stock_actual = 0 WHERE activo = 1 OR activo IS NULL")
    conn.commit()
    conn.close()
    
    return redirect(url_for('index'))

@app.route('/importar_lista', methods=['POST'])
def importar_lista():
    if 'archivo_lista' not in request.files:
        return redirect(url_for('index'))
    file = request.files['archivo_lista']
    if file.filename == '':
        return redirect(url_for('index'))
    
    if file:
        try:
            if file.filename.endswith('.csv'):
                df = pd.read_csv(file)
            else:
                df = pd.read_excel(file)

            conn = conectar()
            cursor = conn.cursor()

            # Iterar sobre las filas importadas
            for _, row in df.iterrows():
                codigo = str(row.get('Codigo', '') if pd.notna(row.get('Codigo')) else '')
                nombre = str(row.get('Producto', ''))
                viscosidad = str(row.get('Viscosidad', ''))
                presentacion = str(row.get('Envase', ''))
                bulto = int(row.get('Bulto_Unidades', 1))
                costo = float(row.get('Costo_Neto', 0.0))
                venta = float(row.get('Venta_Final', 0.0))
                stock = int(row.get('Stock', 0))
                litros_unitarios = extraer_litros_preciso(presentacion + ' ' + nombre)
                tipo_prod = clasificar_tipo_producto(nombre, presentacion)

                cursor.execute('''
                    INSERT INTO productos (codigo_barras, nombre, viscosidad, presentacion, bulto_unidades, litros_unitarios, precio_costo, precio_venta, stock_actual, ubicacion, activo, tipo_producto)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 'MOSTRADOR', 1, ?)
                ''', (codigo, nombre, viscosidad, presentacion, bulto, litros_unitarios, costo, venta, stock, tipo_prod))

            conn.commit()
            conn.close()
        except Exception as e:
            print("Error al importar la lista:", e)

    return redirect(url_for('index'))

@app.route('/dar_de_baja/<int:id>', methods=['POST'])
def dar_de_baja(id):
    conn = conectar()
    cursor = conn.cursor()
    ahora = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    cursor.execute("UPDATE productos SET activo = 0, fecha_baja = ? WHERE id = ?", (ahora, id))
    conn.commit()
    conn.close()
    flash("📦 Producto enviado a la papelera. Se auto-eliminará en 30 días si no se restaura.", "info")
    return redirect(url_for('index'))

@app.route('/restaurar_producto/<int:id>', methods=['POST'])
def restaurar_producto(id):
    conn = conectar()
    cursor = conn.cursor()
    cursor.execute("UPDATE productos SET activo = 1, fecha_baja = NULL WHERE id = ?", (id,))
    conn.commit()
    conn.close()
    flash("🟢 Producto restaurado con éxito al mostrador.", "success")
    return redirect(url_for('index'))

@app.route('/eliminar_definitivo/<int:id>', methods=['POST'])
def eliminar_definitivo(id):
    conn = conectar()
    cursor = conn.cursor()
    try:
        cursor.execute("UPDATE historial_ventas SET producto_id = NULL WHERE producto_id = ?", (id,))
        cursor.execute("DELETE FROM reservas WHERE producto_id = ?", (id,))
        cursor.execute("DELETE FROM productos WHERE id = ?", (id,))
        conn.commit()
        flash("🗑️ Producto eliminado definitivamente de la base de datos.", "info")
    except Exception as e:
        conn.rollback()
        flash(f"❌ Error al eliminar el producto: {str(e)}", "error")
    finally:
        conn.close()
    return redirect(url_for('index'))

@app.route('/vaciar_papelera', methods=['POST'])
def vaciar_papelera():
    conn = conectar()
    cursor = conn.cursor()
    try:
        # 1. Desvincular ventas históricas para no romper reportes ni violar FK
        cursor.execute("""
            UPDATE historial_ventas 
            SET producto_id = NULL 
            WHERE producto_id IN (SELECT id FROM productos WHERE activo = 0)
        """)
        # 2. Eliminar reservas asociadas a los productos dados de baja
        cursor.execute("""
            DELETE FROM reservas 
            WHERE producto_id IN (SELECT id FROM productos WHERE activo = 0)
        """)
        # 3. Eliminar definitivamente los productos inactivos
        cursor.execute("DELETE FROM productos WHERE activo = 0")
        conn.commit()
        
        # 4. Compactar base de datos
        try:
            cursor.execute("VACUUM")
        except Exception:
            pass
            
        flash("🗑️ Papelera vaciada por completo y espacio liberado con éxito.", "success")
    except Exception as e:
        conn.rollback()
        flash(f"❌ Error al vaciar la papelera: {str(e)}", "error")
    finally:
        conn.close()
    return redirect(url_for('index'))

@app.route('/optimizar_base_datos', methods=['POST'])
def optimizar_base_datos():
    try:
        conn = conectar()
        cursor = conn.cursor()
        cursor.execute("PRAGMA optimize")
        cursor.execute("VACUUM")
        conn.close()
        flash("⚡ Base de datos compactada, índices optimizados y espacio liberado al 100%.", "success")
    except Exception as e:
        flash(f"❌ Error al optimizar la base de datos: {str(e)}", "error")
    return redirect(url_for('index'))

@app.route('/reiniciar_ventas', methods=['POST'])
def reiniciar_ventas():
    try:
        conn = conectar()
        cursor = conn.cursor()
        cursor.execute("DELETE FROM historial_ventas")
        conn.commit()
        conn.close()
        flash("✅ Historial de ventas y métricas reiniciados correctamente.", "success")
    except Exception as e:
        flash(f"❌ Error al reiniciar métricas: {str(e)}", "error")
    return redirect(url_for('index'))

@app.route('/anular_remito', methods=['POST'])
def anular_remito():
    fecha_remito = request.form.get('fecha_remito', '').strip()
    if not fecha_remito:
        flash("❌ Fecha de remito inválida.", "error")
        return redirect(url_for('index'))
    try:
        conn = conectar()
        cursor = conn.cursor()
        cursor.execute("SELECT producto_id, cantidad FROM historial_ventas WHERE fecha = ?", (fecha_remito,))
        ventas = cursor.fetchall()
        for v in ventas:
            prod_id, cant = v
            if prod_id:
                cursor.execute("UPDATE productos SET stock_actual = stock_actual + ? WHERE id = ?", (cant, prod_id))
        cursor.execute("DELETE FROM historial_ventas WHERE fecha = ?", (fecha_remito,))
        conn.commit()
        conn.close()
        flash(f"✅ Remito del {fecha_remito} anulado. Se restauró el stock de los productos y se descontaron los litros.", "success")
    except Exception as e:
        flash(f"❌ Error al anular el remito: {str(e)}", "error")
    return redirect(url_for('index'))

# ENDPOINTS DE API JSON PARA REDUCIR CARGA DE RED Y CPU
@app.route('/api/remito_detalle')
def api_remito_detalle():
    """Devuelve el desglose completo de productos de un remito para el modal."""
    fecha = request.args.get('fecha', '').strip()
    if not fecha:
        return jsonify({'success': False, 'message': 'Fecha de remito requerida'}), 400
    try:
        conn = conectar()
        cursor = conn.cursor()
        cursor.execute("""
            SELECT 
                h.id,
                COALESCE(h.nombre_producto, p.nombre, h.marca || ' ' || h.viscosidad) as nombre,
                h.marca,
                h.viscosidad,
                h.presentacion,
                h.cantidad,
                h.litros_totales,
                h.precio_total,
                h.tipo_producto
            FROM historial_ventas h
            LEFT JOIN productos p ON h.producto_id = p.id
            WHERE h.fecha = ?
            ORDER BY h.id ASC
        """, (fecha,))
        rows = cursor.fetchall()
        conn.close()

        items = []
        total_litros_liq = 0.0
        total_kg_grasa = 0.0
        total_monto = 0.0
        for r in rows:
            tipo = r[8] or 'LIQUIDO'
            volumen = float(r[6] or 0.0)
            monto = float(r[7] or 0.0)
            if tipo == 'GRASA':
                total_kg_grasa += volumen
            else:
                total_litros_liq += volumen
            total_monto += monto

            items.append({
                'id': r[0],
                'nombre': r[1],
                'marca': r[2],
                'viscosidad': r[3],
                'presentacion': r[4],
                'cantidad': r[5],
                'volumen': volumen,
                'precio_total': monto,
                'tipo_producto': tipo
            })

        return jsonify({
            'success': True,
            'fecha': fecha,
            'total_items': len(items),
            'total_litros_liq': total_litros_liq,
            'total_kg_grasa': total_kg_grasa,
            'total_litros': total_litros_liq + total_kg_grasa,
            'total_monto': total_monto,
            'items': items
        })
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)}), 500

@app.route('/api/reservas_pendientes')
def api_reservas_pendientes():
    try:
        conn = conectar()
        cursor = conn.cursor()
        cursor.execute('''
            SELECT r.id, p.nombre, p.viscosidad, p.presentacion, r.empleado_nombre, r.cantidad, r.estado, r.fecha, p.stock_actual
            FROM reservas r
            JOIN productos p ON r.producto_id = p.id
            WHERE r.estado = 'PENDIENTE'
            ORDER BY r.fecha DESC
        ''')
        reservas = cursor.fetchall()
        conn.close()
        
        lista_reservas = []
        for r in reservas:
            lista_reservas.append({
                'id': r[0],
                'producto_nombre': r[1],
                'viscosidad': r[2],
                'presentacion': r[3],
                'vendedor': r[4],
                'cantidad': r[5],
                'estado': r[6],
                'fecha': r[7],
                'stock_actual': r[8]
            })
        return jsonify({
            'success': True,
            'reservas': lista_reservas,
            'count': len(lista_reservas)
        })
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/producto/<string:codigo_o_id>')
def api_producto(codigo_o_id):
    try:
        codigo_o_id = codigo_o_id.strip()
        conn = conectar()
        cursor = conn.cursor()
        # Buscar por código de barras exacto, o por ID si coincide
        cursor.execute('''
            SELECT id, codigo_barras, nombre, viscosidad, presentacion, bulto_unidades, precio_venta, stock_actual 
            FROM productos 
            WHERE (codigo_barras = ? OR id = ?) AND (activo = 1 OR activo IS NULL)
        ''', (codigo_o_id, codigo_o_id))
        prod = cursor.fetchone()
        conn.close()
        
        if prod:
            return jsonify({
                'success': True,
                'producto': {
                    'id': prod[0],
                    'codigo': prod[1],
                    'nombre': prod[2],
                    'viscosidad': prod[3],
                    'presentacion': prod[4],
                    'bulto': prod[5],
                    'precio_venta': prod[6],
                    'stock': prod[7]
                }
            })
        return jsonify({'success': False, 'message': 'Lubricante no encontrado'}), 404
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/actualizar_stock_ajax', methods=['POST'])
def api_actualizar_stock_ajax():
    try:
        data = request.get_json()
        id_producto = data.get('id_producto')
        diferencial = data.get('diferencial')
        
        conn = conectar()
        cursor = conn.cursor()
        
        cursor.execute("SELECT stock_actual, nombre FROM productos WHERE id = ?", (id_producto,))
        row = cursor.fetchone()
        if not row:
            conn.close()
            return jsonify({'success': False, 'message': 'Producto no encontrado'}), 404
        
        stock_actual = row[0] if row[0] is not None else 0
        nombre_prod = row[1]
        
        if diferencial is not None:
            nuevo_stock = stock_actual + int(diferencial)
        elif data.get('incrementar_uno', False):
            nuevo_stock = stock_actual + 1
        else:
            nuevo_stock = int(data.get('nuevo_stock', 0))
            
        if nuevo_stock < 0:
            nuevo_stock = 0
            
        cursor.execute("UPDATE productos SET stock_actual = ? WHERE id = ?", (nuevo_stock, id_producto))
        conn.commit()
        conn.close()
        
        return jsonify({
            'success': True,
            'id': id_producto,
            'nombre': nombre_prod,
            'nuevo_stock': nuevo_stock,
            'message': 'Stock actualizado'
        })
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)}), 500

# =========================================================================
# RUTAS DE SINCRONIZACIÓN Y PERFILES (RAMIRO MASTER <-> MOSTRADOR SLAVE)
# =========================================================================

@app.route('/cambiar_modo', methods=['POST'])
def cambiar_modo():
    nuevo_modo = request.form.get('modo', 'mostrador')
    pin = request.form.get('pin', '').strip()
    pin_actual = obtener_pin_admin()
    if nuevo_modo == 'admin':
        if pin and pin == pin_actual:
            session['modo'] = 'admin'
            flash("🔓 Modo Administrador activado.", "success")
        else:
            flash("❌ Contraseña de Administrador incorrecta.", "error")
    else:
        session['modo'] = 'mostrador'
        flash("🖥️ Modo Mostrador (Solo Consulta) activado.", "success")
    return redirect(url_for('index'))

@app.route('/guardar_pin_admin', methods=['POST'])
def guardar_pin_admin():
    nuevo_pin = request.form.get('nuevo_pin', '').strip()
    if len(nuevo_pin) < 3:
        flash("❌ La contraseña debe tener al menos 3 caracteres.", "error")
    else:
        if actualizar_pin_admin(nuevo_pin):
            flash(f"✅ Contraseña de Administrador guardada con éxito: '{nuevo_pin}'", "success")
        else:
            flash("❌ Error al guardar la contraseña en la base de datos.", "error")
    return redirect(url_for('index'))

@app.route('/api/sync/enviar_mostrador', methods=['POST'])
def api_sync_enviar_mostrador():
    """Ramiro envía la base de datos actualizada a la compu de Mostrador en 1 clic."""
    try:
        data = request.get_json() or {}
        ip_mostrador = data.get('ip_mostrador', '').strip()
        if not ip_mostrador:
            return jsonify({'success': False, 'message': 'Debe ingresar la dirección IP de la computadora del Mostrador.'}), 400
        
        # Normalizar IP y puerto
        if not ip_mostrador.startswith('http://') and not ip_mostrador.startswith('https://'):
            ip_mostrador = f"http://{ip_mostrador}"
        if ':5000' not in ip_mostrador and not any(char.isdigit() for char in ip_mostrador.split(':')[-1:] if len(ip_mostrador.split(':')) > 2):
            ip_mostrador = f"{ip_mostrador}:5000"
            
        url_destino = f"{ip_mostrador.rstrip('/')}/api/sync/recibir_db"
        
        # Forzar guardado completo en disco de SQLite antes de enviar
        conn = conectar()
        conn.execute("PRAGMA wal_checkpoint(FULL)")
        conn.close()
        
        db_path = get_db_path()
        with open(db_path, 'rb') as f:
            files = {'db_file': (os.path.basename(db_path), f, 'application/octet-stream')}
            headers = {'X-Sync-Token': SYNC_TOKEN}
            resp = requests.post(url_destino, files=files, headers=headers, timeout=10.0)
            
        if resp.status_code == 200 and resp.json().get('success'):
            return jsonify({'success': True, 'message': '✅ ¡Base de datos y catálogo sincronizados con éxito en el Mostrador!'})
        else:
            msg = resp.json().get('message', f'Error HTTP {resp.status_code}') if resp.status_code == 200 else f"Error HTTP {resp.status_code}"
            return jsonify({'success': False, 'message': f"No se pudo sincronizar: {msg}"}), 400
    except requests.exceptions.ConnectionError:
        return jsonify({'success': False, 'message': '❌ No se pudo conectar con la computadora de Mostrador. Verifique que esté encendida, conectada al mismo Wi-Fi y con el sistema abierto.'}), 500
    except Exception as e:
        return jsonify({'success': False, 'message': f'Error al sincronizar: {str(e)}'}), 500

@app.route('/api/sync/recibir_db', methods=['POST'])
def api_sync_recibir_db():
    """La compu de Mostrador recibe y actualiza la base de datos de forma segura."""
    token = request.headers.get('X-Sync-Token', '')
    if token != SYNC_TOKEN:
        return jsonify({'success': False, 'message': 'Token de sincronización inválido o no autorizado.'}), 401
    
    if 'db_file' not in request.files:
        return jsonify({'success': False, 'message': 'No se recibió ningún archivo de base de datos.'}), 400
        
    archivo = request.files['db_file']
    if not archivo.filename:
        return jsonify({'success': False, 'message': 'Nombre de archivo vacío.'}), 400
        
    try:
        db_path = get_db_path()
        backup_path = db_path + f".backup_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
        if os.path.exists(db_path):
            try:
                shutil.copy2(db_path, backup_path)
            except Exception:
                pass
            
        archivo.save(db_path)
        init_db()
        return jsonify({'success': True, 'message': 'Base de datos recibida y actualizada correctamente.'})
    except Exception as e:
        return jsonify({'success': False, 'message': f'Error al guardar base de datos: {str(e)}'}), 500

@app.route('/sync/descargar_backup')
def sync_descargar_backup():
    """Descarga una copia de la base de datos para llevar en pendrive."""
    conn = conectar()
    conn.execute("PRAGMA wal_checkpoint(FULL)")
    conn.close()
    db_path = get_db_path()
    fecha = datetime.now().strftime("%Y%m%d_%H%M")
    return send_file(db_path, download_name=f'inventario_oleos_{fecha}.db', as_attachment=True)

@app.route('/sync/restaurar_backup', methods=['POST'])
def sync_restaurar_backup():
    """Restaura una copia de inventario.db desde el pendrive."""
    if 'db_file' not in request.files:
        flash("❌ No se seleccionó ningún archivo.", "error")
        return redirect(url_for('index'))
    archivo = request.files['db_file']
    if not archivo.filename or not archivo.filename.endswith('.db'):
        flash("❌ El archivo debe ser una base de datos con extensión .db", "error")
        return redirect(url_for('index'))
    try:
        db_path = get_db_path()
        backup_path = db_path + f".backup_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
        if os.path.exists(db_path):
            try:
                shutil.copy2(db_path, backup_path)
            except Exception:
                pass
        archivo.save(db_path)
        init_db()
        flash("✅ Base de datos restaurada con éxito desde el pendrive.", "success")
    except Exception as e:
        flash(f"❌ Error al restaurar base de datos: {str(e)}", "error")
    return redirect(url_for('index'))

# MANEJO GLOBAL DE ERRORES PARA PREVENIR ERROR 500 Y PANTALLAZOS EN BLANCO
@app.errorhandler(sqlite3.Error)
def handle_db_error(error):
    app.logger.error(f"Error de Base de Datos: {error}")
    if request.path.startswith('/api/'):
        return jsonify({'success': False, 'error': f'Error de base de datos: {error}'}), 500
    
    if request.endpoint == 'index':
        # Evitamos bucles de redirección si el error es en el index
        return render_template(
            'index.html',
            productos=[],
            productos_deposito=[],
            productos_archivados=[],
            reservas_pendientes=[],
            hace_60_dias='',
            total_litros=0.0,
            total_litros_liquidos=0.0,
            total_kg_grasas=0.0,
            total_facturacion=0.0,
            metricas_marca=[],
            metricas_viscosidad=[],
            resumen_mensual=[],
            stock_bajo_count=0,
            agotados_count=0,
            db_error=str(error)
        )
    flash(f"⚠️ Error de base de datos: {error}", "error")
    return redirect(url_for('index'))

@app.errorhandler(Exception)
def handle_generic_error(error):
    app.logger.error(f"Error Inesperado: {error}")
    if request.path.startswith('/api/'):
        return jsonify({'success': False, 'error': 'Error interno del servidor'}), 500
    
    if request.endpoint == 'index':
        return render_template(
            'index.html',
            productos=[],
            productos_deposito=[],
            productos_archivados=[],
            reservas_pendientes=[],
            hace_60_dias='',
            total_litros=0.0,
            total_litros_liquidos=0.0,
            total_kg_grasas=0.0,
            total_facturacion=0.0,
            metricas_marca=[],
            metricas_viscosidad=[],
            resumen_mensual=[],
            stock_bajo_count=0,
            agotados_count=0,
            db_error=str(error)
        )
    flash("⚠️ Ocurrió un error inesperado en el servidor.", "error")
    return redirect(url_for('index'))

if __name__ == '__main__':
    init_db()
    app.run(host='0.0.0.0', port=5000, debug=True)